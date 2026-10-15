"""Loads staff from an Excel file (.xlsx): document number, names, area and position.

The document number is the username and also the first password. A document that already exists
(in the database or earlier in the same file) is never duplicated or modified.
"""

import unicodedata

from openpyxl import load_workbook

from .models import User

# Accepted header names (compared without accents, case or spaces) for each column.
HEADERS = {
    "document": {"documento", "numerodedocumento", "nrodocumento", "cedula", "identificacion", "numerodeidentificacion"},
    "first_name": {"nombres", "nombre", "nombrescompletos", "nombrecompleto"},
    "last_name": {"apellidos", "apellido"},
    "area": {"area"},
    "position": {"cargo"},
}
REQUIRED = ("document", "first_name", "area", "position")
LABELS = {"document": "documento", "first_name": "nombres", "area": "área", "position": "cargo"}


class StaffFileError(Exception):
    """The file cannot be read or does not have the expected columns."""


def _key(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in text.lower() if ch.isalnum())


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():  # Excel stores 1012345678 as 1012345678.0
        value = int(value)
    return " ".join(str(value).split())


def _read_rows(file):
    try:
        sheet = load_workbook(file, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise StaffFileError("No se pudo leer el archivo. Sube un Excel válido (.xlsx).") from exc

    header = next((i for i, row in enumerate(rows) if any(cell not in (None, "") for cell in row)), None)
    if header is None:
        raise StaffFileError("El archivo está vacío.")

    columns = {}
    for index, cell in enumerate(rows[header]):
        for field, names in HEADERS.items():
            if _key(cell) in names and field not in columns:
                columns[field] = index
    missing = [LABELS[f] for f in REQUIRED if f not in columns]
    if missing:
        raise StaffFileError(f"Faltan columnas en el encabezado: {', '.join(missing)}.")
    return header, columns, rows


def import_staff(file) -> dict:
    header, columns, rows = _read_rows(file)

    existing = set(User.objects.values_list("username", flat=True))
    created, skipped, errors = 0, [], []
    for number, row in enumerate(rows[header + 1:], start=header + 2):
        values = {f: _text(row[i]) if i < len(row) else "" for f, i in columns.items()}
        if not any(values.values()):
            continue  # blank line

        missing = [LABELS[f] for f in REQUIRED if not values[f]]
        if missing:
            errors.append(f"Fila {number}: falta {', '.join(missing)}.")
            continue
        document = values["document"]
        if document in existing:
            skipped.append(document)
            continue

        User.objects.create_user(
            username=document,
            password=document,
            first_name=values["first_name"],
            last_name=values.get("last_name", ""),
            role=User.Role.COLLABORATOR,
            area=values["area"],
            position=values["position"],
        )
        existing.add(document)
        created += 1

    return {"created": created, "skipped": len(skipped), "skipped_documents": skipped, "errors": errors}
