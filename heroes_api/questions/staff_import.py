"""Loads evaluation answers from an Excel (.xlsx) or CSV file.

One row = the scores that one collaborator gave to one colleague in one evaluation:

    evaluation | collaborator | personal | <one column per question of that evaluation>

- evaluation: id or exact name of the evaluation. May be left empty when `default_evaluation` is given or
  there is a single evaluation.
- collaborator / personal: the username (document number) of who scores and who is scored.
- one column per question: its code (see the admin, "Criteria"), with a whole score from 1 to 5. A row only
  needs the questions of ITS evaluation; the other columns can stay empty.

The area is not asked: it is the one of the scored person. A collaborator can score a colleague only once per
evaluation: if the pair already has answers in that evaluation, the row is skipped (never duplicated).
"""

import csv
import io
import unicodedata

from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import load_workbook

from users.models import User

from .models import Category, Criterion, Evaluation, Question

# Accepted header names (compared without accents, case, spaces or symbols).
HEADERS = {
    "evaluation": {"evaluation", "evaluacion"},
    "collaborator": {"collaborator", "colaborador", "evaluador", "documentocolaborador"},
    "personal": {"personal", "evaluado", "documentopersonal"},
}


class AnswersFileError(Exception):
    """The file cannot be read or does not have the expected columns."""


def _key(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in text.lower() if ch.isalnum())


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():  # Excel stores 1012345678 as 1012345678.0
        value = int(value)
    return " ".join(str(value).split())


def read_table(file, filename: str = "") -> list[dict[str, str | int]]:
    """The rows of the file as dicts keyed by the normalized header, plus the line number under "_row"."""
    if filename.lower().endswith(".csv"):
        try:
            text = file.read().decode("utf-8-sig")
        except (UnicodeDecodeError, AttributeError) as exc:
            raise AnswersFileError("No se pudo leer el CSV (debe estar en UTF-8).") from exc
        grid = list(csv.reader(io.StringIO(text)))
    else:
        try:
            sheet = load_workbook(file, read_only=True, data_only=True).worksheets[0]
            grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        except Exception as exc:
            raise AnswersFileError("No se pudo leer el archivo. Sube un Excel válido (.xlsx) o un CSV.") from exc

    start = next((i for i, row in enumerate(grid) if any(_text(cell) for cell in row)), None)
    if start is None:
        raise AnswersFileError("El archivo está vacío.")

    keys = [_key(cell) for cell in grid[start]]
    table = []
    for number, row in enumerate(grid[start + 1:], start=start + 2):
        values: dict[str, str | int] = {key: _text(row[i]) if i < len(row) else "" for i, key in enumerate(keys) if key}
        if any(values.values()):
            values["_row"] = number
            table.append(values)
    return table


def _pick(row: dict, field: str) -> str:
    return next((str(row[name]) for name in HEADERS[field] if row.get(name)), "")


def _find_evaluation(raw: str, default):
    if not raw:
        return default or (Evaluation.objects.get() if Evaluation.objects.count() == 1 else None)
    if raw.isdigit():
        by_id = Evaluation.objects.filter(pk=raw).first()
        if by_id is not None:
            return by_id
    return Evaluation.objects.filter(name__iexact=raw).first()


# The questions that make up the original standard form, grouped by category.
STANDARD_FORM = (
    ("capacidad_de_respuesta", "Capacidad de respuesta", (
        ("META", "Demuestra cumplimiento de metas, objetivos y funciones en el área asignada en los tiempos establecidos."),
        ("EFECTIVA", "Mantiene respuestas oportunas en la labor asignada, y gestiona los recursos personales e institucionales para generar una respuesta efectiva."),
        ("PRODUCTIVIDAD", "Demuestra productividad y efectividad en la relación tiempo-tarea, en los diferentes contextos institucionales."),
    )),
    ("apropiacion_del_rol", "Apropiación del rol", (
        ("IMPLEMENTA", "Implementa y comparte sus habilidades y conocimientos con el equipo de trabajo."),
        ("APRENDIZAJE", "Demuestra aprendizaje continuo y capacidad de respuesta al cambio."),
        ("RECEPTIVA", "Mantiene una actitud receptiva y de compromiso ante las diversas situaciones promoviendo el trabajo articulado."),
    )),
    ("comunicacion", "Comunicación", (
        ("RESPETA", "Respeta el conducto regular acatando las directrices y desarrolla una comunicación clara y precisa con los grupos de interés."),
        ("ESCUCHA", "Maneja una escucha activa, con respuestas asertivas a diferentes grupos de interés."),
        ("FOMENTA", "Fomenta un ambiente laboral positivo desde el fortalecimiento de las relaciones interpersonales."),
    )),
    ("compromiso_organizacional", "Compromiso organizacional", (
        ("INTERACCION", "Genera interacción y respuesta asertiva en el equipo de trabajo y con otras áreas manteniendo la imagen de la institución."),
        ("CONOCIMIENTO", "Demuestra conocimiento, adherencia y cumplimiento a las políticas institucionales."),
        ("DISPOSICION", "Demuestra disposición e interés en cooperar y dar respuesta oportuna a necesidades del área y/o institucionales."),
    )),
)


EVALUATION_HEADERS = {"evaluation", "evaluacion", "nombreevaluacion"}
CATEGORY_HEADERS = {"category", "categoria", "nombrecategoria"}
CRITERION_HEADERS = {"question", "pregunta", "criterion", "criterio"}


def import_evaluations(table: list[dict[str, str | int]]) -> dict:
    """Create evaluation forms from rows with evaluation, category and question columns.

    Repeated rows build one evaluation with multiple categories and questions. An
    evaluation already in the database is skipped as a whole to avoid duplicating it.
    """
    if not table:
        raise AnswersFileError("El archivo no contiene filas de evaluaciones.")

    headers = set(table[0])
    if not headers.intersection(EVALUATION_HEADERS):
        raise AnswersFileError("Falta la columna evaluation (nombre de la evaluación).")
    if not headers.intersection(CATEGORY_HEADERS):
        raise AnswersFileError("Falta la columna category (nombre de la categoría).")
    if not headers.intersection(CRITERION_HEADERS):
        raise AnswersFileError("Falta la columna question (texto de la pregunta).")

    def value(row, choices):
        return next((str(row[key]).strip() for key in choices if row.get(key)), "")

    forms = {}
    errors = []
    for row in table:
        line = int(row["_row"])
        evaluation_name = value(row, EVALUATION_HEADERS)
        category_name = value(row, CATEGORY_HEADERS)
        question_text = value(row, CRITERION_HEADERS)
        if not evaluation_name or not category_name or not question_text:
            errors.append(f"Fila {line}: debe incluir evaluación, categoría y pregunta.")
            continue
        if len(evaluation_name) > 150:
            errors.append(f"Fila {line}: el nombre de la evaluación supera 150 caracteres.")
            continue
        if len(category_name) > 100:
            errors.append(f"Fila {line}: el nombre de la categoría supera 100 caracteres.")
            continue

        evaluation_key = evaluation_name.casefold()
        form = forms.setdefault(evaluation_key, {"name": evaluation_name, "categories": {}})
        category_key = category_name.casefold()
        category = form["categories"].setdefault(category_key, {"name": category_name, "questions": []})
        if question_text.casefold() not in {item.casefold() for item in category["questions"]}:
            category["questions"].append(question_text)

    created = skipped = criteria_created = 0
    created_names = []
    for form in forms.values():
        name = form["name"]
        if Evaluation.objects.filter(name__iexact=name).exists():
            skipped += 1
            continue

        evaluation = Evaluation.objects.create(name=name)
        created += 1
        created_names.append(name)
        question_index = 0
        for category_order, category_data in enumerate(form["categories"].values(), start=1):
            base = slugify(category_data["name"])[:35] or "category"
            code = f"{base}-{evaluation.pk}-{category_order}"[:50]
            category = Category.objects.create(
                evaluation=evaluation,
                code=code,
                name=category_data["name"],
                order=category_order,
            )
            for question_order, question_text in enumerate(category_data["questions"], start=1):
                question_index += 1
                Criterion.objects.create(
                    category=category,
                    code=f"q-{evaluation.pk}-{question_index}",
                    name=question_text,
                    order=question_order,
                )
                criteria_created += 1

    return {"created": created, "skipped": skipped, "criteria_created": criteria_created,
            "evaluations_created": created_names, "errors": errors}


def import_answers(table: list[dict[str, str | int]], default_evaluation: Evaluation | None = None) -> dict:
    """Creates the answers of every row. Returns {"created", "skipped", "errors"} (errors say the row number)."""
    created, skipped, errors = 0, 0, []
    forms = {}  # evaluation id -> {code: (header key, category code)}

    for row in table:
        number = int(row["_row"])
        evaluation = _find_evaluation(_pick(row, "evaluation"), default_evaluation)
        if evaluation is None:
            errors.append(f"Fila {number}: falta la evaluación, o no existe (usa su id o su nombre exacto).")
            continue

        if evaluation.pk not in forms:
            forms[evaluation.pk] = {
                code: (_key(code), category)
                for code, category in Criterion.active(evaluation).values_list("code", "category__code")
            }
        form = forms[evaluation.pk]
        if not form:
            errors.append(f"Fila {number}: la evaluación «{evaluation.name}» no tiene preguntas.")
            continue

        collaborator = User.objects.filter(username=_pick(row, "collaborator"), role=User.Role.COLLABORATOR).first()
        personal = User.objects.filter(username=_pick(row, "personal"), role=User.Role.COLLABORATOR).first()
        if collaborator is None or personal is None:
            errors.append(f"Fila {number}: no existe el colaborador o la persona evaluada (usa su número de documento).")
            continue
        if collaborator.pk == personal.pk:
            errors.append(f"Fila {number}: un colaborador no puede evaluarse a sí mismo.")
            continue

        scores, problems = {}, []
        for code, (key, _category) in form.items():
            raw = str(row.get(key, ""))
            if not raw.isdigit() or not 1 <= int(raw) <= 5:
                problems.append(code)
            else:
                scores[code] = int(raw)
        if problems:
            errors.append(f"Fila {number}: falta una nota de 1 a 5 en: {', '.join(problems)}.")
            continue

        if Question.objects.filter(evaluation=evaluation, collaborator=collaborator, personal=personal).exists():
            skipped += 1
            continue

        now = timezone.now()
        with transaction.atomic():
            Question.objects.bulk_create(
                Question(
                    evaluation=evaluation,
                    collaborator=collaborator,
                    personal=personal,
                    area=personal.area,
                    category=form[code][1],
                    criterion=code,
                    score=score,
                    reviewed_at=now,
                )
                for code, score in scores.items()
            )
        created += 1

    return {"created": created, "skipped": skipped, "errors": errors}
